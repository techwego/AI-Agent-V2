import io
from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy import func, or_
from datetime import datetime, timedelta
import re
from pydantic import BaseModel

from backend.database.db import get_db
from backend.database.models import User, Book, Department, Upload, AdminLog, ConversationHistory, LibraryConfig
from backend.auth.auth_middleware import require_admin

router = APIRouter(prefix="/api/admin", tags=["analytics", "logs"])

def extract_book_title_from_query(raw_query: str) -> str:
    """
    Extracts and normalizes clean book titles from full conversational queries.
    E.g. 'i want this book ponniyin selvan' -> 'Ponniyin Selvan'
         'where is goodnight moon?' -> 'Goodnight Moon'
         'do you have the book alchemist available?' -> 'Alchemist'
    """
    if not raw_query:
        return ""
    q = raw_query.strip()
    
    # 1. Remove common conversational prefixes/intents
    prefix_patterns = [
        r'^(?:hey\s+sam|sam|hi|hello|please|kindly)[,\s]+',
        r'^(?:can\s+you|could\s+you|please)?\s*(?:find|search(?:\s+for)?|locate|show(?:\s+me)?|give(?:\s+me)?|tell(?:\s+me)?(?:\s+about)?|get(?:\s+me)?)\s+(?:the\s+)?(?:book\s+)?(?:called\s+|named\s+|titled\s+)?',
        r'^(?:i\s+want|i\s+need|i\s+am\s+looking\s+for|i\'m\s+looking\s+for|do\s+you\s+have|is\s+there|have\s+you\s+got)\s+(?:a\s+|an\s+|the\s+)?(?:copy\s+of\s+)?(?:book\s+|novel\s+)?(?:called\s+|named\s+|titled\s+)?(?:this\s+book\s+|that\s+book\s+)?',
        r'^(?:where\s+is|where\s+can\s+i\s+find|where\s+are)\s+(?:the\s+)?(?:book\s+)?(?:called\s+|named\s+)?',
        r'^(?:book\s+named|book\s+called|book\s+titled|book\s+on|novel\s+called)\s+',
        r'^(?:is\s+)?(?:the\s+)?(?:book\s+)?'
    ]
    
    changed = True
    while changed:
        old_q = q
        for pat in prefix_patterns:
            q = re.sub(pat, '', q, flags=re.IGNORECASE).strip()
        if q == old_q:
            changed = False
            
    # 2. Remove common conversational suffixes
    suffix_patterns = [
        r'\s+(?:available|in\s+stock|present|available\s+now|in\s+the\s+library|in\s+library|here)\??$',
        r'\s+(?:book|novel|copies|copy)\??$',
        r'\s+(?:please|kindly|sam)\??$',
        r'[?.!]+$'
    ]
    changed = True
    while changed:
        old_q = q
        for pat in suffix_patterns:
            q = re.sub(pat, '', q, flags=re.IGNORECASE).strip()
        if q == old_q:
            changed = False
            
    # Strip quotes and double spaces
    q = q.strip('"\' ')
    q = re.sub(r'\s+', ' ', q).strip()
    
    if q and len(q) >= 2:
        general_stops = {"how many books", "total books", "library hours", "rules", "who are you", "what is your name", "hello", "hi", "help"}
        if q.lower() in general_stops:
            return ""
        return q.title()
    return ""

from typing import Optional

@router.get("/analytics")
def get_analytics(start_date: Optional[str] = None, end_date: Optional[str] = None, db: Session = Depends(get_db), current_user: User = Depends(require_admin)):
    total_books = db.query(func.count(Book.id)).scalar() or 0
    total_users = db.query(func.count(User.id)).scalar() or 0
    total_departments = db.query(func.count(Department.id)).scalar() or 0
    
    today = datetime.utcnow().date()
    today_queries = db.query(func.count(ConversationHistory.id)).filter(func.date(ConversationHistory.created_at) == today).scalar() or 0
    
    total_uploads = db.query(func.count(Upload.id)).scalar() or 0
    active_users = db.query(func.count(User.id)).filter(User.is_active == True).scalar() or 0
    
    # Missing Books Analysis (Unfound Queries)
    unanswered_query = db.query(ConversationHistory).filter(
        or_(
            ConversationHistory.response.ilike("%could not find%"),
            ConversationHistory.response.ilike("%not found%"),
            ConversationHistory.response.ilike("%not in our%"),
            ConversationHistory.response.ilike("%no records found%")
        )
    )

    if start_date:
        try:
            start_dt = datetime.strptime(start_date, "%Y-%m-%d")
            unanswered_query = unanswered_query.filter(ConversationHistory.created_at >= start_dt)
        except ValueError:
            pass
            
    if end_date:
        try:
            end_dt = datetime.strptime(end_date, "%Y-%m-%d") + timedelta(days=1)
            unanswered_query = unanswered_query.filter(ConversationHistory.created_at < end_dt)
        except ValueError:
            pass

    unanswered_logs = unanswered_query.order_by(ConversationHistory.created_at.desc()).limit(1000).all()
    
    missing_books = {}
    for log in unanswered_logs:
        clean_title = extract_book_title_from_query(log.query)
        if clean_title and len(clean_title) >= 2:
            missing_books[clean_title] = missing_books.get(clean_title, 0) + 1
            
    # Sort top missing books
    top_missing = sorted([{"title": k, "searches": v} for k, v in missing_books.items()], key=lambda x: x["searches"], reverse=True)[:15]
    
    # 7-day query trend
    trend_data = []
    for i in range(6, -1, -1):
        target_date = today - timedelta(days=i)
        count = db.query(func.count(ConversationHistory.id)).filter(func.date(ConversationHistory.created_at) == target_date).scalar() or 0
        trend_data.append({"date": target_date.strftime("%b %d"), "queries": count})
    
    # Recent activity
    recent_logs = db.query(AdminLog).order_by(AdminLog.created_at.desc()).limit(10).all()
    recent_activity = []
    for log in recent_logs:
        admin_user = db.query(User).filter(User.id == log.admin_id).first()
        recent_activity.append({
            "id": log.id,
            "action": log.action,
            "details": log.details,
            "admin_name": admin_user.username if admin_user else "Unknown",
            "created_at": log.created_at
        })
    
    return {
        "total_books": total_books,
        "total_users": total_users,
        "total_departments": total_departments,
        "today_queries": today_queries,
        "total_uploads": total_uploads,
        "active_users": active_users,
        "recent_activity": recent_activity,
        "top_missing_books": top_missing,
        "trend_data": trend_data
    }

@router.get("/logs")
def get_admin_logs(db: Session = Depends(get_db), current_user: User = Depends(require_admin)):
    try:
        cutoff_24h = datetime.utcnow() - timedelta(hours=24)
        db.query(AdminLog).filter(AdminLog.created_at < cutoff_24h).delete()
        db.commit()
    except Exception:
        db.rollback()
    return db.query(AdminLog).order_by(AdminLog.created_at.desc()).limit(100).all()

@router.get("/chat-logs")
def get_chat_logs(db: Session = Depends(get_db), current_user: User = Depends(require_admin)):
    try:
        cutoff_24h = datetime.utcnow() - timedelta(hours=24)
        db.query(ConversationHistory).filter(ConversationHistory.created_at < cutoff_24h).delete()
        db.commit()
    except Exception:
        db.rollback()
    logs = db.query(ConversationHistory).order_by(ConversationHistory.created_at.desc()).limit(100).all()
    return logs

class DeleteLogsRequest(BaseModel):
    before_date: str # YYYY-MM-DD

@router.delete("/chat-logs")
def delete_chat_logs(req: DeleteLogsRequest, db: Session = Depends(get_db), current_user: User = Depends(require_admin)):
    try:
        cutoff_date = datetime.strptime(req.before_date, "%Y-%m-%d")
        deleted_count = db.query(ConversationHistory).filter(ConversationHistory.created_at < cutoff_date).delete()
        db.commit()
        return {"message": f"Successfully deleted {deleted_count} logs older than {req.before_date}.", "deleted": deleted_count}
    except Exception as e:
        db.rollback()
        raise HTTPException(status_code=400, detail=str(e))

@router.get("/analytics/export-missing-books-excel")
def export_missing_books_excel(db: Session = Depends(get_db), current_user: User = Depends(require_admin)):
    """Generates and streams an enterprise formatted .xlsx Excel workbook of Missing Book Demands."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    # Get Library details
    cfg = db.query(LibraryConfig).first()
    college_name = cfg.college_name if cfg and cfg.college_name else "University College of Engineering"
    library_name = cfg.library_name if cfg and cfg.library_name else "Central Library"

    # Query unanswered / missing logs
    unanswered_logs = db.query(ConversationHistory).filter(
        or_(
            ConversationHistory.response.ilike("%could not find%"),
            ConversationHistory.response.ilike("%not found%"),
            ConversationHistory.response.ilike("%not in our%"),
            ConversationHistory.response.ilike("%no records found%")
        )
    ).order_by(ConversationHistory.created_at.desc()).limit(2000).all()

    missing_map = {}
    for log in unanswered_logs:
        clean_title = extract_book_title_from_query(log.query)
        if clean_title and len(clean_title) >= 2:
            if clean_title not in missing_map:
                missing_map[clean_title] = {
                    "count": 0,
                    "first_date": log.created_at,
                    "last_date": log.created_at,
                    "sample_query": log.query
                }
            missing_map[clean_title]["count"] += 1
            if log.created_at < missing_map[clean_title]["first_date"]:
                missing_map[clean_title]["first_date"] = log.created_at
            if log.created_at > missing_map[clean_title]["last_date"]:
                missing_map[clean_title]["last_date"] = log.created_at

    sorted_missing = sorted(
        [{"title": k, **v} for k, v in missing_map.items()],
        key=lambda x: x["count"],
        reverse=True
    )

    # Build Excel Workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Missing Book Demands"
    ws.views.sheetView[0].showGridLines = True

    # Styling Palette
    navy_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    indigo_fill = PatternFill(start_color="3B82F6", end_color="3B82F6", fill_type="solid")
    light_blue_fill = PatternFill(start_color="EFF6FF", end_color="EFF6FF", fill_type="solid")
    alt_row_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    
    title_font = Font(name="Segoe UI", size=15, bold=True, color="FFFFFF")
    subtitle_font = Font(name="Segoe UI", size=10, italic=True, color="E0E7FF")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    bold_font = Font(name="Segoe UI", size=10, bold=True, color="0F172A")
    regular_font = Font(name="Segoe UI", size=10, color="1E293B")
    
    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

    # Title Banner Block
    ws.merge_cells("A1:F1")
    title_cell = ws["A1"]
    title_cell.value = f"📚 {library_name.upper()} — MISSING BOOK DEMAND REPORT"
    title_cell.font = title_font
    title_cell.fill = navy_fill
    title_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 36

    ws.merge_cells("A2:F2")
    sub_cell = ws["A2"]
    sub_cell.value = f"Institution: {college_name}  |  Generated On: {datetime.now().strftime('%Y-%m-%d %I:%M %p')}  |  Total Demanded Titles: {len(sorted_missing)}"
    sub_cell.font = subtitle_font
    sub_cell.fill = navy_fill
    sub_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[2].height = 20

    # Header Row
    headers = [
        ("Rank", 8),
        ("Demanded Book Title / Subject", 40),
        ("Total Requests", 16),
        ("First Requested", 20),
        ("Latest Request", 20),
        ("Procurement Priority", 22)
    ]
    
    ws.row_dimensions[4].height = 26
    for col_idx, (h_name, width) in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col_idx, value=h_name)
        cell.font = header_font
        cell.fill = indigo_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = width

    # Data Rows
    for idx, item in enumerate(sorted_missing, 1):
        row_num = 4 + idx
        ws.row_dimensions[row_num].height = 22
        
        # Priority level calculation
        requests = item["count"]
        if requests >= 5:
            priority = "🔴 CRITICAL / HIGH DEMAND"
        elif requests >= 2:
            priority = "🟡 MEDIUM DEMAND"
        else:
            priority = "🟢 LOW DEMAND"

        first_str = item["first_date"].strftime("%Y-%m-%d %H:%M") if item["first_date"] else "N/A"
        last_str = item["last_date"].strftime("%Y-%m-%d %H:%M") if item["last_date"] else "N/A"

        row_values = [
            idx,
            item["title"],
            requests,
            first_str,
            last_str,
            priority
        ]

        for col_idx, val in enumerate(row_values, 1):
            cell = ws.cell(row=row_num, column=col_idx, value=val)
            cell.border = thin_border
            
            if col_idx == 1 or col_idx == 3:
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.font = bold_font
            elif col_idx == 2:
                cell.alignment = Alignment(horizontal="left", vertical="center")
                cell.font = bold_font
            elif col_idx == 6:
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.font = bold_font
            else:
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.font = regular_font

            if idx % 2 == 0:
                cell.fill = alt_row_fill

    # Save to BytesIO
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"missing_books_demand_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )

@router.get("/analytics/export-demand-books-excel")
def export_demand_books_excel(db: Session = Depends(get_db), current_user: User = Depends(require_admin)):
    """Exports books with 10+ student requests (high-demand titles not in library catalog) as a styled Excel report."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    cfg = db.query(LibraryConfig).first()
    college_name = cfg.college_name if cfg and cfg.college_name else "University College of Engineering"
    library_name = cfg.library_name if cfg and cfg.library_name else "Central Library"

    unanswered_logs = db.query(ConversationHistory).filter(
        or_(
            ConversationHistory.response.ilike("%could not find%"),
            ConversationHistory.response.ilike("%not found%"),
            ConversationHistory.response.ilike("%not in our%"),
            ConversationHistory.response.ilike("%no records found%")
        )
    ).order_by(ConversationHistory.created_at.desc()).limit(5000).all()

    missing_map = {}
    for log in unanswered_logs:
        clean_title = extract_book_title_from_query(log.query)
        if clean_title and len(clean_title) >= 2:
            if clean_title not in missing_map:
                missing_map[clean_title] = {
                    "count": 0,
                    "first_date": log.created_at,
                    "last_date": log.created_at,
                    "sample_query": log.query
                }
            missing_map[clean_title]["count"] += 1
            if log.created_at < missing_map[clean_title]["first_date"]:
                missing_map[clean_title]["first_date"] = log.created_at
            if log.created_at > missing_map[clean_title]["last_date"]:
                missing_map[clean_title]["last_date"] = log.created_at

    # Filter only books with 10+ requests (High Demand threshold)
    demand_books = sorted(
        [{"title": k, **v} for k, v in missing_map.items() if v["count"] >= 10],
        key=lambda x: x["count"],
        reverse=True
    )

    # Build Excel Workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "High Demand Books"
    ws.views.sheetView[0].showGridLines = True

    # Styling Palette
    deep_red_fill = PatternFill(start_color="991B1B", end_color="991B1B", fill_type="solid")
    red_fill = PatternFill(start_color="DC2626", end_color="DC2626", fill_type="solid")
    rose_light = PatternFill(start_color="FFF1F2", end_color="FFF1F2", fill_type="solid")
    alt_row_fill = PatternFill(start_color="FEF2F2", end_color="FEF2F2", fill_type="solid")

    title_font = Font(name="Segoe UI", size=15, bold=True, color="FFFFFF")
    subtitle_font = Font(name="Segoe UI", size=10, italic=True, color="FECACA")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    bold_font = Font(name="Segoe UI", size=10, bold=True, color="0F172A")
    regular_font = Font(name="Segoe UI", size=10, color="1E293B")

    thin_border = Border(
        left=Side(style='thin', color='FECDD3'),
        right=Side(style='thin', color='FECDD3'),
        top=Side(style='thin', color='FECDD3'),
        bottom=Side(style='thin', color='FECDD3')
    )

    # Title Banner
    ws.merge_cells("A1:G1")
    title_cell = ws["A1"]
    title_cell.value = f"🔥 {library_name.upper()} — HIGH DEMAND BOOKS REPORT (10+ REQUESTS)"
    title_cell.font = title_font
    title_cell.fill = deep_red_fill
    title_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 36

    ws.merge_cells("A2:G2")
    sub_cell = ws["A2"]
    sub_cell.value = f"Institution: {college_name}  |  Generated: {datetime.now().strftime('%Y-%m-%d %I:%M %p')}  |  High Demand Titles: {len(demand_books)}  |  Threshold: ≥10 student requests"
    sub_cell.font = subtitle_font
    sub_cell.fill = deep_red_fill
    sub_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[2].height = 20

    # Header Row
    headers = [
        ("Rank", 8),
        ("Book Title / Subject", 42),
        ("Total Student Requests", 22),
        ("First Requested", 20),
        ("Latest Request", 20),
        ("Demand Level", 20),
        ("Recommended Action", 28)
    ]

    ws.row_dimensions[4].height = 26
    for col_idx, (h_name, width) in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col_idx, value=h_name)
        cell.font = header_font
        cell.fill = red_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = width

    # Data Rows
    for idx, item in enumerate(demand_books, 1):
        row_num = 4 + idx
        ws.row_dimensions[row_num].height = 22

        requests = item["count"]
        if requests >= 50:
            demand_level = "🔴 EXTREMELY HIGH"
            action = "URGENT: Immediate Procurement"
        elif requests >= 25:
            demand_level = "🟠 VERY HIGH"
            action = "Priority Purchase Required"
        elif requests >= 15:
            demand_level = "🟡 HIGH"
            action = "Schedule for Next Purchase Cycle"
        else:
            demand_level = "🟢 MODERATE"
            action = "Add to Procurement Wishlist"

        first_str = item["first_date"].strftime("%Y-%m-%d %H:%M") if item["first_date"] else "N/A"
        last_str = item["last_date"].strftime("%Y-%m-%d %H:%M") if item["last_date"] else "N/A"

        row_values = [idx, item["title"], requests, first_str, last_str, demand_level, action]

        for col_idx, val in enumerate(row_values, 1):
            cell = ws.cell(row=row_num, column=col_idx, value=val)
            cell.border = thin_border
            if col_idx in (1, 3):
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.font = bold_font
            elif col_idx == 2:
                cell.alignment = Alignment(horizontal="left", vertical="center")
                cell.font = bold_font
            elif col_idx in (6, 7):
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.font = bold_font
            else:
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.font = regular_font
            if idx % 2 == 0:
                cell.fill = alt_row_fill

    # Summary Row
    summary_row = 4 + len(demand_books) + 2
    ws.merge_cells(f"A{summary_row}:G{summary_row}")
    summary_cell = ws.cell(row=summary_row, column=1)
    summary_cell.value = f"📊 Total High-Demand Titles: {len(demand_books)}  |  These books have been requested 10 or more times by students but are not currently in the library catalog."
    summary_cell.font = Font(name="Segoe UI", size=10, italic=True, color="991B1B")
    summary_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"high_demand_books_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )
